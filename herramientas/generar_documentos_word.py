#!/usr/bin/env python3
"""Genera los 3 documentos Word de Trafiqueishon con el formato de las plantillas PushProMaster.

Cada documento se construye abriendo su plantilla (así hereda estilos, tipografías, márgenes,
numeración y pie de página), vaciando el cuerpo y escribiendo el nuevo contenido.

Uso:
    pip install python-docx
    python3 generar_documentos_word.py [carpeta]      # por defecto: ../documentos_word

La carpeta debe contener las 3 plantillas PushProMaster*.docx; los resultados se guardan ahí.
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# ---------------------------------------------------------------- datos
TITULO = "Trafiqueishon - Simulador de Tráfico Urbano"
UNIVERSIDAD = "Universidad Nacional de Juliaca"
CURSO = "Dinámica de Sistemas"
DOCENTE = "Tipo Mamani Noe Wilber"
FECHA = "23 de setiembre"
AUTORES = [
    "Jean Piere Emanuel Gonzales Yerba",
    "Luis Fernando Yujra Ramirez",
    "Beto Cuquin Hanco Pomari",
    "Josue Daniel Taza Condori",
    "Eberth Paulo Mamani Flores",
]

PROYECTO = Path(__file__).resolve().parent.parent
JAVA = PROYECTO / "src" / "main" / "java" / "trafiqueishon"
CLASES = [
    ("LuisYMain", JAVA / "LuisYMain.java"),
    ("LuisYVentanaPrincipal", JAVA / "ui" / "LuisYVentanaPrincipal.java"),
    ("EberthMMotorSimulacion", JAVA / "simulacion" / "EberthMMotorSimulacion.java"),
    ("PieroGGeoJsonParser", JAVA / "geo" / "PieroGGeoJsonParser.java"),
]
PLACEHOLDER = "[INSERTAR CAPTURA AQUÍ]"

# plantilla -> (fragmento del nombre, archivo de salida)
PLANTILLAS = {
    "ejecucion": ("ejecutar", "Trafiqueishon_Manual_Ejecucion.docx"),
    "codigo": ("CODIGO", "Trafiqueishon_Codigo_Fuente.docx"),
    "usuario": ("Manual de Uso", "Trafiqueishon_Manual_Usuario.docx"),
}


# ---------------------------------------------------------------- plantillas
def abrir_plantilla(carpeta, clave):
    fragmento = PLANTILLAS[clave][0]
    hallados = [p for p in carpeta.glob("PushProMaster*.docx") if fragmento in p.name]
    if not hallados:
        sys.exit(f"No encuentro la plantilla PushProMaster ({fragmento}) en {carpeta}")
    doc = Document(hallados[0])
    body = doc.element.body
    for hijo in list(body):
        if hijo.tag != qn("w:sectPr"):
            body.remove(hijo)
    # las capturas de la plantilla ya no se usan: se sueltan para no inflar el archivo
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith("/image"):
            del doc.part.rels[rid]
    return doc


def guardar(doc, carpeta, clave):
    destino = carpeta / PLANTILLAS[clave][1]
    doc.save(destino)
    print(f"  ✔ {destino}")


# ---------------------------------------------------------------- utilidades
def fuente(run, nombre=None, tam=None, negrita=None, subrayado=None, cursiva=None):
    if nombre:
        run.font.name = nombre
        rf = run._element.get_or_add_rPr().get_or_add_rFonts()
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), nombre)
    if tam:
        run.font.size = Pt(tam)
    if negrita is not None:
        run.bold = negrita
    if subrayado is not None:
        run.underline = subrayado
    if cursiva is not None:
        run.italic = cursiva


def parrafo(doc_o_celda, texto="", estilo=None, alinear=None, nombre=None, tam=None,
            negrita=None, subrayado=None, despues=None):
    p = doc_o_celda.add_paragraph(style=estilo) if estilo else doc_o_celda.add_paragraph()
    if alinear is not None:
        p.alignment = alinear
    if despues is not None:
        p.paragraph_format.space_after = Pt(despues)
    if texto:
        fuente(p.add_run(texto), nombre, tam, negrita, subrayado)
    return p


def salto_pagina(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def numerar(p, num_id, nivel=0):
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl")
    il.set(qn("w:val"), str(nivel))
    ni = OxmlElement("w:numId")
    ni.set(qn("w:val"), str(num_id))
    numPr.append(il)
    numPr.append(ni)
    pPr.append(numPr)


def tabla_cuadricula(doc, filas, columnas):
    t = doc.add_table(rows=filas, cols=columnas)
    t.style = doc.styles["Table Grid"]
    return t


def celda_texto(celda, texto, **kw):
    p = celda.paragraphs[0]
    if kw.get("alinear") is not None:
        p.alignment = kw["alinear"]
    fuente(p.add_run(texto), kw.get("nombre"), kw.get("tam"), kw.get("negrita"))
    return p


def portada(doc, subtitulo, nombre=None, estilo_titulo=None):
    """Portada con los mismos estilos de la plantilla: todo centrado."""
    centro = WD_ALIGN_PARAGRAPH.CENTER
    for _ in range(3):
        parrafo(doc, "", alinear=centro)
    parrafo(doc, UNIVERSIDAD, alinear=centro, nombre=nombre, tam=16, negrita=True)
    parrafo(doc, "")
    if estilo_titulo:
        t = doc.add_paragraph(TITULO, style=estilo_titulo)
        t.alignment = centro
    else:
        parrafo(doc, TITULO, alinear=centro, nombre=nombre, tam=15, negrita=True)
    parrafo(doc, subtitulo, alinear=centro, nombre=nombre, tam=14)
    parrafo(doc, "")
    parrafo(doc, "")
    for etiqueta, valor in (("Curso", CURSO), ("Docente", DOCENTE), ("Fecha", FECHA)):
        p = parrafo(doc, "", alinear=centro, despues=4)
        fuente(p.add_run(f"{etiqueta}: "), nombre, 12, True)
        fuente(p.add_run(valor), nombre, 12)
    parrafo(doc, "")
    parrafo(doc, "Autores", alinear=centro, nombre=nombre, tam=12, negrita=True, despues=4)
    for a in AUTORES:
        parrafo(doc, a, alinear=centro, nombre=nombre, tam=12, despues=2)
    salto_pagina(doc)


def captura(doc, nombre=None):
    parrafo(doc, PLACEHOLDER, alinear=WD_ALIGN_PARAGRAPH.CENTER, nombre=nombre, negrita=True)


# ---------------------------------------------------------------- ARCHIVO 1: Manual de Ejecución
# Plantilla: A4, Times New Roman, título 14 pt centrado, listas numeradas, tabla Problema/Solución.
def manual_ejecucion(carpeta):
    doc = abrir_plantilla(carpeta, "ejecucion")
    TNR = "Times New Roman"
    portada(doc, "Manual de Ejecución", TNR)

    parrafo(doc, f"Para ejecutar “{TITULO.split(' - ')[0]}”", alinear=WD_ALIGN_PARAGRAPH.CENTER,
            nombre=TNR, tam=14)

    def seccion(texto, nivel=0):
        p = parrafo(doc, texto, estilo="List Paragraph", nombre=TNR)
        numerar(p, 10, nivel)

    def vineta(texto):
        p = parrafo(doc, texto, estilo="List Paragraph", nombre=TNR)
        numerar(p, 7)

    def consola(lineas):
        t = tabla_cuadricula(doc, 1, 1)
        celda = t.rows[0].cells[0]
        for i, linea in enumerate(lineas):
            p = celda.paragraphs[0] if i == 0 else celda.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            fuente(p.add_run(linea), "Consolas", 9.5)
        parrafo(doc, "", nombre=TNR)

    seccion("Requisitos Previos")
    vineta("Tener instalado Java (JDK) 21 o superior (solo para ejecutar desde el código fuente).")
    vineta("Tener instalado Apache Maven 3.8 o superior (solo para ejecutar desde el código fuente).")
    vineta("Memoria RAM de al menos 2 GiB (para el mapa de Lima se recomiendan 4 GiB o más).")
    vineta("Contar con la carpeta mapas/<ciudad>/ con los archivos calles.geojson y puntos.json.")
    parrafo(doc, "Puedes verificar las versiones instaladas con:", nombre=TNR)
    consola(["java -version", "mvn -version"])

    seccion("Ejecución desde el código fuente (.java)")
    seccion("Abrir la carpeta del proyecto", 1)
    vineta("Abre una terminal y entra a la carpeta que contiene el archivo pom.xml.")
    seccion("Probar el proyecto", 1)
    vineta("Ejecuta las pruebas automáticas para confirmar que todo compila correctamente:")
    consola(["cd \"generador de rutas\"", "mvn test"])
    seccion("Ejecutar la aplicación desde la terminal", 1)
    consola(["mvn javafx:run                  # Juliaca / Puno (768 MiB de heap)",
             "mvn javafx:run -Dapp.xmx=2g     # Lima u otros mapas masivos"])
    captura(doc, TNR)

    seccion("Ejecución desde archivo ejecutable")
    seccion("Generar el ejecutable", 1)
    vineta("El paquete incluye su propio JRE, por lo que no hace falta instalar Java para usarlo.")
    consola(["mvn -Pinstalador clean package -DskipTests"])
    seccion("Ejecutar", 1)
    vineta("Abre la carpeta target/Trafiqueishon/ y haz doble clic sobre el archivo Trafiqueishon.")
    vineta("En Linux, si no abre, concede permiso de ejecución:")
    consola(["chmod +x target/Trafiqueishon/Trafiqueishon", "./target/Trafiqueishon/Trafiqueishon"])
    p = parrafo(doc, "", nombre=TNR)
    fuente(p.add_run("Importante: "), TNR, negrita=True)
    fuente(p.add_run("La primera vez, el sistema operativo puede mostrar una advertencia de seguridad "
                     "porque el archivo no está firmado digitalmente. Confirma que confías en el "
                     "archivo para continuar."), TNR)
    captura(doc, TNR)

    seccion("Posibles advertencias o errores comunes")
    errores = [
        ("Problema", "Solución"),
        ("java: command not found", "Instala el JDK 21 y verifica con java -version."),
        ("invalid target release: 21", "El JDK es anterior a la versión 21: actualízalo o configura JAVA_HOME."),
        ("mvn: command not found", "Instala Maven y verifica con mvn -version."),
        ("JavaFX runtime components are missing",
         "Usa mvn javafx:run o el ejecutable; no lances directamente la clase de la ventana."),
        ("OutOfMemoryError (mapa de Lima)", "Ejecuta con más memoria: mvn javafx:run -Dapp.xmx=2g."),
        ("Mapa vacío o aviso de calles.geojson faltante",
         "Coloca el archivo en mapas/<ciudad>/calles.geojson."),
        ("Pantalla negra o parpadeo", "JavaFX usará el pipeline por software; actualiza los drivers de video."),
        ("Permission denied al ejecutar", "Concede permiso de ejecución: chmod +x Trafiqueishon."),
    ]
    t = tabla_cuadricula(doc, len(errores), 2)
    for i, (a, b) in enumerate(errores):
        celda_texto(t.rows[i].cells[0], a, nombre=TNR)
        celda_texto(t.rows[i].cells[1], b, nombre=TNR)
    parrafo(doc, "", nombre=TNR)

    seccion("Consejos")
    vineta("Ejecuta mvn test antes de exponer: confirma que el proyecto está sano.")
    vineta("En equipos con 2 GiB de RAM usa el heap por defecto (768 MiB).")
    vineta("El mapa de Lima debe ir fuera del jar (pesa cientos de MB); se acepta calles.geojson.gz.")
    vineta("Las rutas guardadas viven en ~/.trafiqueishon/: respáldalas antes de reinstalar.")
    vineta("Tab oculta el panel lateral y F11 activa el modo presentación.")
    guardar(doc, carpeta, "ejecucion")


# ---------------------------------------------------------------- ARCHIVO 2: Código Fuente
# Plantilla: A4, título «CÓDIGO FUENTE» 15 pt negrita subrayado, código en tabla con Consolas 9.5 pt.
def codigo_fuente(carpeta):
    doc = abrir_plantilla(carpeta, "codigo")
    centro = WD_ALIGN_PARAGRAPH.CENTER
    portada(doc, "Código Fuente")

    for i, (nombre, ruta) in enumerate(CLASES):
        if i:
            salto_pagina(doc)
        parrafo(doc, "CÓDIGO FUENTE", alinear=centro, tam=15, negrita=True, subrayado=True, despues=0)
        parrafo(doc, f"Clase {nombre}", alinear=centro, tam=12, negrita=True)
        if not ruta.exists():
            print(f"  ⚠ falta {ruta}")
            parrafo(doc, f"[No se encontró {ruta.name}]")
            continue
        lineas = ruta.read_text(encoding="utf-8").splitlines()
        t = tabla_cuadricula(doc, 2, 1)
        celda_texto(t.rows[0].cells[0], f"{ruta.relative_to(PROYECTO)}  ({len(lineas)} líneas)",
                    alinear=centro, tam=15, negrita=True)
        celda = t.rows[1].cells[0]
        for k, linea in enumerate(lineas):
            p = celda.paragraphs[0] if k == 0 else celda.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            fuente(p.add_run(linea.replace("\t", "    ")), "Consolas", 9.5)
    guardar(doc, carpeta, "codigo")


# ---------------------------------------------------------------- ARCHIVO 3: Manual de Usuario
# Plantilla: Carta, estilo «Title» centrado con filete, «Heading 1» numerado a mano,
# texto Normal y viñetas con guion dentro del mismo párrafo.
def manual_usuario(carpeta):
    doc = abrir_plantilla(carpeta, "usuario")
    centro = WD_ALIGN_PARAGRAPH.CENTER
    portada(doc, "Manual de Usuario", estilo_titulo="Title")

    t = doc.add_paragraph("Manual de Uso - Trafiqueishon", style="Title")
    t.alignment = centro

    def h1(texto):
        doc.add_paragraph(texto, style="Heading 1")

    def texto(contenido):
        doc.add_paragraph(contenido)

    def guiones(items):
        p = doc.add_paragraph()
        for i, it in enumerate(items):
            r = p.add_run(("\n" if i else "") + "- " + it)

    h1("1. Introducción")
    texto("Trafiqueishon es una aplicación de escritorio para visualizar la red vial de Juliaca, Puno "
          "o Lima, editar puntos de interés, definir rutas de micros, simular sus salidas programadas "
          "con eventos de tráfico y exportar las rutas en distintos formatos. Desde la Dinámica de "
          "Sistemas, modela el transporte urbano como un sistema donde la demanda de pasajeros, la "
          "flota, el tráfico, los semáforos y los eventos se retroalimentan en el tiempo.")
    guiones([
        "Cargar el mapa de la ciudad y sus puntos de interés",
        "Crear rutas de micros manualmente o con «Autogenerar Rutas»",
        "Simular con hora punta, lluvia, trancaderas, accidentes y cierres",
        "Ver estadísticas en vivo y el mapa de calor de demanda",
        "Exportar a PDF, GeoJSON, GPX, CSV y TXT",
    ])

    h1("2. Interfaz Principal")
    texto("La ventana principal incluye los siguientes elementos:")
    guiones([
        "Barra superior: selector de Ciudad, botones Iniciar/Pausar (▶) y Detener (⏹), reloj de la "
        "simulación y botón «Solo mapa (Tab)»",
        "Mapa central: calles, puntos de interés, rutas, vehículos y eventos, con la leyenda flotante "
        "en la esquina inferior derecha",
        "Panel lateral con secciones plegables: Mapas, Rutas, Ida / Vuelta, Puntos de Interés, "
        "Configuración de Ciudad, Simulación, Eventos y Exportar",
    ])
    captura(doc)
    texto("Atajos de teclado:")
    guiones([
        "Tab: oculta o muestra el panel lateral",
        "F11: activa el modo presentación; Esc: sale de él",
        "Espacio: inicia, pausa o continúa la simulación (no actúa mientras se escribe en un campo)",
    ])

    h1("3. Flujo de Uso")
    pasos = [
        ("1. Elige la ciudad.",
         "Selecciona Juliaca, Puno o Lima en el selector «Ciudad». La aplicación recarga calles, "
         "puntos, rutas guardadas y eventos.", "Selector de ciudad y mapa cargado"),
        ("2. Revisa los puntos de interés.",
         "En «Puntos de Interés» busca, agrega, edita o importa (JSON/Excel) los puntos que generan "
         "demanda: educación, salud, mercados, plazas y transporte.", "Sección Puntos de Interés"),
        ("3. Crea las rutas.",
         "Pulsa «🤖 Autogenerar Rutas» para que el sistema proponga rutas según la demanda, o define "
         "las tuyas con al menos dos paraderos. Marca la casilla de cada micro para verla en el mapa.",
         "Diálogo Autogenerar Rutas"),
        ("4. Revisa la ida y la vuelta.",
         "Selecciona una micro y pulsa «⇄ Activar»: el mapa muestra solo su ida (azul), su vuelta "
         "(roja), flechas de sentido y el panel de paraderos.", "Modo Ida / Vuelta"),
        ("5. (Opcional) Agrega eventos.",
         "En «Eventos» elige lluvia, trancadera, accidente o cierre y haz clic sobre la vía. Los "
         "cierres y accidentes recalculan desvíos con A* sin detener la simulación.", "Colocación de un evento"),
        ("6. Ejecuta la simulación.",
         "En «Simulación» fija la hora inicial y marca las micros. Pulsa ▶ (o Espacio) para iniciar "
         "y de nuevo para pausar; activa «📊 Estadísticas sobre el mapa» para ver abordajes y "
         "ocupación en vivo. Al pulsar ⏹ aparece un resumen con abordajes, ocupación media, "
         "incidencias y tiempo simulado.", "Simulación en marcha"),
    ]
    for titulo_paso, cuerpo, _ in pasos:
        p = doc.add_paragraph()
        p.add_run(titulo_paso + " ").bold = True
        p.add_run(cuerpo)
        captura(doc)
    texto("El reloj simulado aumenta el tráfico entre 07:00–09:00, 12:00–14:00 y 17:00–20:00, y la "
          "demanda (mañana 2.5x, mediodía 1.8x, tarde 2.8x, noche 0.2x).")

    h1("4. Leyenda")
    texto("Pulsa «🎨 Leyenda» en la esquina inferior derecha del mapa para desplegarla.")
    captura(doc)
    for cat, items in (
        ("Pasajeros", ["Azul: esperando en la parada", "Violeta: transbordo (espera o camina a su 2.º micro)",
                       "Naranja: caminando hacia la parada"]),
        ("Rutas", ["Línea sólida: ida", "Línea punteada: vuelta"]),
        ("Vehículos", ["Color de su ruta: en marcha", "Gris: detenido en semáforo o parada",
                       "Rojo carmesí: micro con sobrecupo"]),
        ("Semáforos", ["Verde: paso", "Rojo: alto", "Amarillo: precaución"]),
        ("Eventos", ["Triángulo rojo: accidente", "Humo gris: trancadera sobre la vía",
                     "Punto naranja: marcador de trancadera / semáforo roto",
                     "Zona azul con gotas: lluvia", "Punto dorado: gasolinera (detiene la unidad 3 min simulados)",
                     "Rojo translúcido: vía cerrada / nodo bloqueado"]),
        ("Otros", ["Policía de tránsito", "Zona sin cobertura: demanda sin parada a ≤ 300 m"]),
    ):
        p = doc.add_paragraph()
        p.add_run(cat).bold = True
        p.add_run("".join("\n- " + it for it in items))

    h1("5. Exportación")
    texto("Las opciones están en la sección «⬇ Exportar» del panel lateral:")
    guiones([
        "PDF · Vista múltiple: una lámina con todas las micros marcadas (ida continua, vuelta "
        "discontinua), lista de rutas y kilómetros sobre el croquis",
        "PDF · Vista individual: ida (azul), vuelta (roja) y paraderos de la micro seleccionada",
        "Otros formatos: GeoJSON, GPX, CSV, TXT o PDF con una lámina por micro, eligiendo rutas y sentidos",
        "Exportar Rutas con Eventos: guarda las rutas junto con los eventos colocados",
    ])
    texto("Pasos: 1. Ve a «Exportar». 2. Elige el formato. 3. Marca las rutas y sentidos "
          "(Marcar visibles / Marcar todas / Desmarcar). 4. Escoge carpeta y nombre; la barra de "
          "estado confirma la ruta del archivo.")
    captura(doc)
    guardar(doc, carpeta, "usuario")


# ---------------------------------------------------------------- main
def main():
    carpeta = Path(sys.argv[1]) if len(sys.argv) > 1 else PROYECTO / "documentos_word"
    print(f"Plantillas y resultados en {carpeta} …")
    manual_ejecucion(carpeta)
    codigo_fuente(carpeta)
    manual_usuario(carpeta)
    print("Listo.")


if __name__ == "__main__":
    main()
